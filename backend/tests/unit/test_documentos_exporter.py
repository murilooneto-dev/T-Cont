import io
from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.application.use_cases.exportacao_use_cases import LinhaExportacao
from app.infrastructure.spreadsheet.documentos_exporter import (
    XLSX_MEDIA_TYPE,
    gerar_planilha_documentos,
)

CABECALHO = [
    "Arquivo", "Data do pagamento", "Valor", "Tipo", "Pagador (nome)",
    "Pagador (CPF/CNPJ)", "Recebedor (nome)", "Recebedor (CPF/CNPJ)", "Banco",
    "Débito (código)", "Débito (descrição)", "Crédito (código)", "Crédito (descrição)",
    "Origem",
]


def _linha(**sobrescritas) -> LinhaExportacao:
    campos = dict(
        arquivo="comprovante.pdf", data_pagamento=date(2026, 9, 18),
        valor=Decimal("150.00"), tipo="PIX", pagador_nome="Tesserato",
        pagador_documento="01234567000199", recebedor_nome="Energisa",
        recebedor_documento="11222333000199", banco_nome="Itau",
        debito_codigo="1", debito_descricao="Energia",
        credito_codigo="2", credito_descricao="Banco Itau", origem="REGRA",
    )
    campos.update(sobrescritas)
    return LinhaExportacao(**campos)


def _abrir(conteudo: bytes):
    return load_workbook(io.BytesIO(conteudo))


def test_media_type_e_o_do_excel():
    assert XLSX_MEDIA_TYPE == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def test_sem_linhas_gera_apenas_o_cabecalho():
    aba = _abrir(gerar_planilha_documentos([]))["Documentos"]

    assert [celula.value for celula in aba[1]] == CABECALHO
    assert aba.max_row == 1


def test_aba_unica_chamada_documentos():
    workbook = _abrir(gerar_planilha_documentos([_linha()]))

    assert workbook.sheetnames == ["Documentos"]


def test_linha_de_dados_na_ordem_das_colunas():
    aba = _abrir(gerar_planilha_documentos([_linha()]))["Documentos"]

    valores = [celula.value for celula in aba[2]]

    assert valores == [
        "comprovante.pdf", datetime(2026, 9, 18), 150, "PIX", "Tesserato",
        "01234567000199", "Energisa", "11222333000199", "Itau", "1", "Energia",
        "2", "Banco Itau", "REGRA",
    ]


def test_data_e_valor_sao_tipos_nativos_e_documentos_sao_texto():
    aba = _abrir(gerar_planilha_documentos([_linha()]))["Documentos"]

    assert aba["B2"].data_type == "d"
    assert aba["C2"].data_type == "n"
    assert aba["F2"].data_type == "s"
    assert aba["F2"].value == "01234567000199"


def test_campos_none_viram_celula_vazia():
    linha = _linha(
        data_pagamento=None, valor=None, tipo=None, pagador_nome=None,
        pagador_documento=None, recebedor_nome=None, recebedor_documento=None,
        banco_nome=None, debito_codigo=None, debito_descricao=None,
        credito_codigo=None, credito_descricao=None, origem=None,
    )
    aba = _abrir(gerar_planilha_documentos([linha]))["Documentos"]

    valores = [celula.value for celula in aba[2]]

    assert valores == ["comprovante.pdf"] + [None] * 13


@pytest.mark.parametrize("texto", ["=1+1", "+cmd", "-2+3", "@SUM(A1)", "#N/A"])
def test_texto_livre_nunca_vira_formula_ou_erro(texto):
    aba = _abrir(gerar_planilha_documentos([_linha(recebedor_nome=texto)]))["Documentos"]

    assert aba["G2"].value == texto
    assert aba["G2"].data_type == "s"
