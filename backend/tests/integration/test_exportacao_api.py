import functools
import io
import re
import time

import fitz
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.deps import get_db, get_storage
from app.infrastructure.db.base import Base
from app.infrastructure.db import models  # noqa: F401
from app.infrastructure.db.session import get_engine
from app.infrastructure.storage.file_storage import LocalFileStorageService
from app.infrastructure.workers.lote_worker import processar_lote_em_background
from app.main import app

CABECALHO = [
    "Arquivo", "Data do pagamento", "Valor", "Tipo", "Pagador (nome)",
    "Pagador (CPF/CNPJ)", "Recebedor (nome)", "Recebedor (CPF/CNPJ)", "Banco",
    "Conta (código)", "Conta (descrição)", "Origem",
]


@pytest.fixture
def client(tmp_path, monkeypatch):
    engine = get_engine("sqlite:///:memory:", poolclass=StaticPool)
    Base.metadata.create_all(engine)
    TestSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)

    def override_get_db():
        session = TestSessionLocal()
        try:
            yield session
            session.commit()
        finally:
            session.close()

    def override_get_storage():
        return LocalFileStorageService(tmp_path)

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_storage] = override_get_storage
    monkeypatch.setattr("app.core.config.settings.storage_root", str(tmp_path))
    monkeypatch.setattr(
        "app.api.routers.lotes_router.processar_lote_em_background",
        functools.partial(processar_lote_em_background, session_factory=TestSessionLocal),
    )
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture
def empresa_id(client):
    response = client.post(
        "/empresas",
        json={"razao_social": "Tesserato", "nome_fantasia": None, "cnpj": "12345678000199"},
    )
    return response.json()["id"]


def _pdf(texto: str) -> bytes:
    documento = fitz.open()
    pagina = documento.new_page()
    pagina.insert_text((72, 72), texto)
    conteudo = documento.tobytes()
    documento.close()
    return conteudo


def _processar_e_aguardar(client, empresa_id):
    lote_id = client.post(f"/empresas/{empresa_id}/documentos/processar").json()["id"]
    status_final = None
    for _ in range(20):
        status_final = client.get(f"/lotes/{lote_id}").json()
        if status_final["status"] != "EM_ANDAMENTO":
            break
        time.sleep(0.5)
    assert status_final["status"] == "CONCLUIDO"


def _abrir(response):
    return load_workbook(io.BytesIO(response.content))


def test_empresa_inexistente_retorna_404(client):
    response = client.get("/empresas/999/documentos/exportar")

    assert response.status_code == 404


def test_empresa_sem_documentos_retorna_xlsx_so_com_cabecalho(client, empresa_id):
    response = client.get(f"/empresas/{empresa_id}/documentos/exportar")

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert re.fullmatch(
        r'attachment; filename="comprovantes_12345678000199_\d{4}-\d{2}-\d{2}\.xlsx"',
        response.headers["content-disposition"],
    )
    aba = _abrir(response)["Documentos"]
    assert [celula.value for celula in aba[1]] == CABECALHO
    assert aba.max_row == 1


def test_documento_pendente_nao_aparece_na_planilha(client, empresa_id):
    client.post(
        f"/empresas/{empresa_id}/documentos",
        files={"arquivos": ("a.pdf", b"conteudo", "application/pdf")},
    )

    response = client.get(f"/empresas/{empresa_id}/documentos/exportar")

    assert _abrir(response)["Documentos"].max_row == 1


def test_documento_processado_e_classificado_aparece_com_todos_os_dados(client, empresa_id):
    plano_id = client.post(
        f"/empresas/{empresa_id}/planos-contas", json={"nome": "Plano"}
    ).json()["id"]
    conta_id = client.post(
        f"/planos-contas/{plano_id}/contas",
        json={
            "codigo": "1", "descricao": "Energia", "natureza": "DESPESA",
            "conta_analitica": True, "conta_pai_id": None,
        },
    ).json()["id"]
    client.post(
        f"/empresas/{empresa_id}/regras",
        json={
            "conta_id": conta_id, "lado_alvo": "RECEBEDOR",
            "documento_fiscal": "11222333000199", "tipo_documento": None,
            "valor_min": None, "valor_max": None, "palavra_chave_nome": None,
        },
    )
    client.post(
        f"/empresas/{empresa_id}/documentos",
        files={
            "arquivos": (
                "comprovante.pdf",
                _pdf("Favorecido: Energisa Distribuidora\nCNPJ: 11.222.333/0001-99\nValor: R$ 150,00"),
                "application/pdf",
            )
        },
    )
    _processar_e_aguardar(client, empresa_id)

    response = client.get(f"/empresas/{empresa_id}/documentos/exportar")

    aba = _abrir(response)["Documentos"]
    assert aba.max_row == 2
    linha = {titulo: celula.value for titulo, celula in zip(CABECALHO, aba[2])}
    assert linha["Arquivo"] == "comprovante.pdf"
    assert linha["Valor"] == 150
    assert linha["Recebedor (CPF/CNPJ)"] == "11222333000199"
    assert linha["Conta (código)"] == "1"
    assert linha["Conta (descrição)"] == "Energia"
    assert linha["Origem"] == "REGRA"


def test_cors_expoe_content_disposition_para_o_frontend(client, empresa_id):
    response = client.get(
        f"/empresas/{empresa_id}/documentos/exportar",
        headers={"Origin": "http://localhost:5173"},
    )

    assert "content-disposition" in response.headers["access-control-expose-headers"].lower()
