import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.application.use_cases.exportacao_use_cases import LinhaExportacao

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_NOME_ABA = "Documentos"
_COLUNAS = [
    ("Arquivo", 30),
    ("Data do pagamento", 18),
    ("Valor", 14),
    ("Tipo", 10),
    ("Pagador (nome)", 28),
    ("Pagador (CPF/CNPJ)", 20),
    ("Recebedor (nome)", 28),
    ("Recebedor (CPF/CNPJ)", 20),
    ("Banco", 22),
    ("Conta (código)", 14),
    ("Conta (descrição)", 30),
    ("Origem", 10),
]
_COLUNA_DATA = 2
_COLUNA_VALOR = 3


def gerar_planilha_documentos(linhas: list[LinhaExportacao]) -> bytes:
    workbook = Workbook()
    aba = workbook.active
    aba.title = _NOME_ABA

    for indice, (titulo, largura) in enumerate(_COLUNAS, start=1):
        celula = aba.cell(row=1, column=indice, value=titulo)
        celula.font = Font(bold=True)
        aba.column_dimensions[get_column_letter(indice)].width = largura

    for numero_linha, linha in enumerate(linhas, start=2):
        valores = [
            linha.arquivo, linha.data_pagamento, linha.valor, linha.tipo,
            linha.pagador_nome, linha.pagador_documento, linha.recebedor_nome,
            linha.recebedor_documento, linha.banco_nome, linha.conta_codigo,
            linha.conta_descricao, linha.origem,
        ]
        for numero_coluna, valor in enumerate(valores, start=1):
            celula = aba.cell(row=numero_linha, column=numero_coluna, value=valor)
            if isinstance(valor, str):
                celula.data_type = "s"
        aba.cell(row=numero_linha, column=_COLUNA_DATA).number_format = "DD/MM/YYYY"
        aba.cell(row=numero_linha, column=_COLUNA_VALOR).number_format = "#,##0.00"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
